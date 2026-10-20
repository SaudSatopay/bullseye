#!/usr/bin/env python3
"""pptx_kit.py · fill an organiser's PPTX template natively (route C), then prove the file is clean.

CLI
    python pptx_kit.py inspect template.pptx [--slide N] [--json]
        Slide size, layouts, fonts, document properties, and every shape on every slide
        (id, name, type, x / y / w / h in inches, text). Plan the fill from this, never by eye.
    python pptx_kit.py check deck.pptx [--max-mb 10] [--expect-slides N] [--require-notes]
                                       [--require-embedded-fonts]
        ERROR on negative extents (cx/cy < 0: PowerPoint opens the file with a Repair prompt that
        python-pptx never notices), WARN on shapes off the slide and empty placeholders (they show
        "Click to add text" in edit view), report fonts used vs embedded. Stdlib only.

Library (put this file next to your fill script, or add its folder to sys.path)
    from pptx_kit import *
    prs = open_pptx("template.pptx")
    s = prs.slides[2]
    clear_text(shape_by_id(s, 3), keep=1)       # blank the prompt bullets, keep the heading line
    text_box(s, 0.6, 1.5, 6.0, 1.0, "Recall **90%** at top 5%", font="Segoe UI", size=18, color="1F2A33")
    card(s, 0.6, 2.8, 4.0, 1.6, "01 · SENSE", "Reads every fridge **every 60 s**.", accent="2F7DD1")
    x = chip(s, 0.6, 4.6, "IMPLEMENTED"); chip(s, x, 4.6, "SIMULATED")
    line(s, 5.0, 3.0, 2.0, 3.0, color="E0563B", arrow="end")    # right-to-left arrow, valid extents
    stat(s, 7.0, 1.5, 2.6, "90%", "RECALL @ TOP 5% · SIMULATED")
    picture_fit(s, crop_image("console.png", (0, 0.08, 1, 0.62)), 6.5, 1.5, 6.0, 3.4, mode="cover")
    set_notes(s, "Say the number, then the baseline it beats.")
    prs.save("deck.pptx")    # then: python pptx_kit.py check deck.pptx --require-notes
All positions and sizes are inches; colours are hex strings ("1F2A33") or RGBColor.
Needs python-pptx (pip install python-pptx); crop_image and picture_fit also need Pillow.
"""
from __future__ import annotations

import json
import os
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

EMU_PER_IN = 914400
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
SHAPE_TAGS = {f"{{{NS['p']}}}{t}" for t in ("sp", "pic", "cxnSp", "grpSp", "graphicFrame")}
STATUS_STYLE = {"IMPLEMENTED": "filled", "PROTOTYPED": "outline", "SIMULATED": "dashed", "FUTURE": "muted"}
MONO_FONT = "Cascadia Mono"   # system mono on Windows 11; use "Courier New" for any machine


# ---------------------------------------------------------------- library: basics
def _pptx():
    try:
        import pptx  # noqa: F401
    except ImportError:
        print("pptx_kit: python-pptx is missing. Install it with:  pip install python-pptx", file=sys.stderr)
        sys.exit(1)
    import pptx
    return pptx


def emu(inches: float) -> int:
    return int(round(inches * EMU_PER_IN))


def rgb(color):
    from pptx.dml.color import RGBColor
    if color is None or isinstance(color, RGBColor):
        return color
    return RGBColor.from_string(str(color).lstrip("#").upper())


def open_pptx(path: str):
    _pptx()
    from pptx import Presentation
    return Presentation(path)


def shape_by_id(slide, shape_id: int):
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            return sh
    raise KeyError(f"no shape with id {shape_id} on this slide (run: pptx_kit.py inspect)")


def delete_shape(slide, shape_or_id) -> None:
    sh = shape_by_id(slide, shape_or_id) if isinstance(shape_or_id, int) else shape_or_id
    sh._element.getparent().remove(sh._element)


def send_to_back(slide, shape, index: int = 2) -> None:
    """Move a shape behind the others (index 2 sits right after the tree's own properties), so the
    template's art stays on top of a photo panel you add."""
    tree = slide.shapes._spTree
    tree.remove(shape._element)
    tree.insert(index, shape._element)


def set_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def alt_text(shape, text: str) -> None:
    shape._element.xpath("./*[1]/p:cNvPr")[0].set("descr", text)


def _normalize(shape):
    """Never write a negative cx/cy: move the origin, flip instead (valid OOXML, no Repair prompt)."""
    xfrm = shape._element.find(".//a:xfrm", NS)
    if xfrm is None:
        xfrm = shape._element.find(".//p:xfrm", NS)
    if xfrm is None:
        return shape
    off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
    if off is None or ext is None:
        return shape
    for o_attr, e_attr, flip in (("x", "cx", "flipH"), ("y", "cy", "flipV")):
        size = int(ext.get(e_attr))
        if size < 0:
            off.set(o_attr, str(int(off.get(o_attr)) + size))
            ext.set(e_attr, str(-size))
            xfrm.set(flip, "0" if xfrm.get(flip) == "1" else "1")
    return shape


# ---------------------------------------------------------------- library: text
def _add_runs(paragraph, text: str, font, size, color, bold) -> None:
    from pptx.util import Pt
    for i, chunk in enumerate(text.split("**")):
        if not chunk:
            continue
        run = paragraph.add_run()
        run.text = chunk
        f = run.font
        if font:
            f.name = font
        if size:
            f.size = Pt(size)
        if color is not None:
            f.color.rgb = rgb(color)
        f.bold = bool(bold) or (i % 2 == 1)


def _write(text_frame, text, *, font, size, color, bold, align, line_spacing, space_after) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    aligns = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
    paras = text if isinstance(text, (list, tuple)) else str(text).split("\n")
    for i, para in enumerate(paras):
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.alignment = aligns[align]
        if line_spacing:
            p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after)
        _add_runs(p, para, font, size, color, bold)


def clear_text(shape, keep: int = 0) -> None:
    """Remove every paragraph after the first `keep` (keep=0 also empties the first one)."""
    tf = shape.text_frame
    for p in list(tf.paragraphs)[max(keep, 1):]:
        p._p.getparent().remove(p._p)
    if keep == 0:
        for r in list(tf.paragraphs[0].runs):
            r._r.getparent().remove(r._r)


def replace_text(shape, text: str) -> None:
    """Swap a placeholder's words but keep the template's font, size and colour (first run wins)."""
    tf = shape.text_frame
    p = tf.paragraphs[0]
    runs = p.runs
    if not runs:
        p.add_run().text = text
    else:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    for extra in list(tf.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)


def text_box(slide, x, y, w, h, text, *, font=None, size=14, color="111111", bold=False, align="left",
             valign="top", line_spacing=1.1, space_after=0, margin=0.0, wrap=True):
    """Text box; `text` is a string ("\\n" starts a paragraph, **bold** spans) or a list of paragraphs."""
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    if w <= 0 or h <= 0:
        raise ValueError(f"text_box needs a positive size, got w={w} h={h}")
    tb = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = emu(margin)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[valign]
    _write(tf, text, font=font, size=size, color=color, bold=bold, align=align, line_spacing=line_spacing, space_after=space_after)
    return tb


# ---------------------------------------------------------------- library: shapes
def rect(slide, x, y, w, h, *, fill=None, line=None, line_w=0.75, radius=None, dash=None):
    """Rectangle (rounded when radius is a 0-0.5 fraction of the short side). Negative sizes are flipped, never written."""
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
    _normalize(sh)
    if radius:
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(line_w)
        if dash:
            sh.line.dash_style = getattr(MSO_LINE_DASH_STYLE, dash.upper())
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return sh


def line(slide, x1, y1, x2, y2, *, color="111111", width=1.25, arrow=None, dash=None):
    """Straight connector from (x1, y1) to (x2, y2) in any direction; arrow = 'end', 'start' or 'both'.
    The box is stored from its top-left corner with flipH/flipV, so cx/cy are never negative."""
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.oxml.ns import qn
    from pptx.util import Pt
    cx = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
    _normalize(cx)
    cx.line.color.rgb = rgb(color)
    cx.line.width = Pt(width)
    if dash:
        cx.line.dash_style = getattr(MSO_LINE_DASH_STYLE, dash.upper())
    if arrow:
        ln = cx.line._get_or_add_ln()
        ends = {"end": ["tailEnd"], "start": ["headEnd"], "both": ["headEnd", "tailEnd"]}[arrow]
        for tag in ends:   # DrawingML: headEnd sits at the start point, tailEnd at the end point
            el = ln.find(qn(f"a:{tag}"))
            if el is None:
                el = ln.makeelement(qn(f"a:{tag}"), {})
                ln.append(el)
            el.set("type", "triangle")
            el.set("w", "med")
            el.set("len", "med")
    return cx


def card(slide, x, y, w, h, title, body, *, fill="FFFFFF", line="D5DCE3", accent=None, radius=0.05,
         title_font=None, body_font=None, title_size=14, body_size=11, title_color="111111",
         body_color="3A4550", pad=0.16):
    """Rounded card: optional accent bar on the left, a title line, a body (both accept **bold**)."""
    box = rect(slide, x, y, w, h, fill=fill, line=line, radius=radius)
    inset = pad
    if accent:
        rect(slide, x, y + 0.02, 0.06, h - 0.04, fill=accent)
        inset = pad + 0.06
    from pptx.util import Pt
    tb = text_box(slide, x + inset, y + pad, w - inset - pad, h - 2 * pad, [f"**{title}**" if title else "", body],
                  font=body_font, size=body_size, color=body_color, space_after=4)
    for r in tb.text_frame.paragraphs[0].runs:   # the title line gets its own face, size and colour
        r.font.size = Pt(title_size)
        if title_font:
            r.font.name = title_font
        r.font.color.rgb = rgb(title_color)
    return box


def chip(slide, x, y, text, *, style=None, color="111111", paper="FFFFFF", font=MONO_FONT, size=9, h=0.28):
    """Mono label chip. style: filled, outline, dashed, muted; status words pick their own:
    IMPLEMENTED filled · PROTOTYPED outline · SIMULATED dashed · FUTURE muted. Returns the x after it."""
    style = style or STATUS_STYLE.get(text.upper(), "outline")
    w = 0.24 + len(text) * size * 0.0088          # mono advance ~0.6 em
    if style == "filled":
        rect(slide, x, y, w, h, fill=color, line=color)
        ink = paper
    elif style == "muted":
        rect(slide, x, y, w, h, line="9AA4AE")
        ink = "7A848E"
    else:
        rect(slide, x, y, w, h, line=color, dash="dash" if style == "dashed" else None)
        ink = color
    text_box(slide, x, y, w, h, text, font=font, size=size, color=ink, bold=True, align="center", valign="middle")
    return x + w + 0.1


def stat(slide, x, y, w, number, label, *, num_font=None, num_size=40, num_color="111111",
         label_font=MONO_FONT, label_size=10, label_color="5A6570", rule=None):
    """Big number over a mono label, with an optional rule on top (the deck's stat block)."""
    top = y
    if rule:
        rect(slide, x, y, w, 0.04, fill=rule)
        top += 0.1
    num_h = num_size / 72 * 1.15
    text_box(slide, x, top, w, num_h, number, font=num_font, size=num_size, color=num_color, bold=True, line_spacing=0.9)
    text_box(slide, x, top + num_h + 0.04, w, 0.6, label, font=label_font, size=label_size, color=label_color, line_spacing=1.15)


def _image_size(path):
    try:
        from PIL import Image
    except ImportError:
        print("pptx_kit: Pillow is missing. Install it with:  pip install pillow", file=sys.stderr)
        sys.exit(1)
    with Image.open(path) as im:
        return im.size


def picture_fit(slide, path, x, y, w, h, *, mode="contain"):
    """Place an image inside a box: contain = whole image, letterboxed and centred; cover = fill the
    box and crop the overflow evenly (PowerPoint crop, the file keeps every pixel)."""
    iw, ih = _image_size(path)
    box_ratio, img_ratio = w / h, iw / ih
    if mode == "contain":
        if img_ratio > box_ratio:
            pw, ph = w, w / img_ratio
        else:
            pw, ph = h * img_ratio, h
        return slide.shapes.add_picture(str(path), emu(x + (w - pw) / 2), emu(y + (h - ph) / 2), emu(pw), emu(ph))
    if mode != "cover":
        raise ValueError("mode must be 'contain' or 'cover'")
    pic = slide.shapes.add_picture(str(path), emu(x), emu(y), emu(w), emu(h))
    if img_ratio > box_ratio:
        cut = (1 - box_ratio / img_ratio) / 2
        pic.crop_left = pic.crop_right = cut
    else:
        cut = (1 - img_ratio / box_ratio) / 2
        pic.crop_top = pic.crop_bottom = cut
    return pic


def crop_image(src, box, out=None) -> str:
    """Crop a screenshot. box = (left, top, right, bottom) in pixels, or all values <= 1 as fractions.
    Returns the path written (default: <name>-crop.png next to the source)."""
    try:
        from PIL import Image
    except ImportError:
        print("pptx_kit: Pillow is missing. Install it with:  pip install pillow", file=sys.stderr)
        sys.exit(1)
    with Image.open(src) as im:
        if all(0 <= v <= 1 for v in box):
            box = (box[0] * im.width, box[1] * im.height, box[2] * im.width, box[3] * im.height)
        box = tuple(int(round(v)) for v in box)
        if box[2] <= box[0] or box[3] <= box[1]:
            raise ValueError(f"empty crop box {box}")
        out = out or os.path.splitext(str(src))[0] + "-crop.png"
        im.crop(box).save(out)
    return out


# ---------------------------------------------------------------- CLI: inspect
def _kind(sh) -> str:
    if sh.is_placeholder:
        return f"PH:{str(sh.placeholder_format.type).split('.')[-1].split(' ')[0]}"
    try:
        return str(sh.shape_type).split(".")[-1].split(" ")[0]
    except Exception:
        return "UNKNOWN"


def _shape_rows(shapes, depth=0):
    for sh in shapes:
        text = ""
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            text = re.sub(r"\s+", " ", sh.text_frame.text).strip()
        def inch(v):
            return None if v is None else round(v / EMU_PER_IN, 2)
        yield {"depth": depth, "id": sh.shape_id, "name": sh.name, "type": _kind(sh), "x": inch(sh.left), "y": inch(sh.top),
               "w": inch(sh.width), "h": inch(sh.height), "text": text}
        if _kind(sh).startswith("GROUP"):
            yield from _shape_rows(sh.shapes, depth + 1)


def cmd_inspect(path: str, only=None, as_json=False) -> int:
    _pptx()
    from pptx import Presentation
    prs = Presentation(path)
    sw, sh_ = prs.slide_width, prs.slide_height
    facts = scan_package(path)
    data = {
        "file": os.path.basename(path), "slides": len(prs.slides),
        "size_in": [round(sw / EMU_PER_IN, 3), round(sh_ / EMU_PER_IN, 3)], "size_emu": [sw, sh_],
        "properties": {"title": prs.core_properties.title, "author": prs.core_properties.author,
                       "last_modified_by": prs.core_properties.last_modified_by},
        "fonts_used": facts["fonts_used"], "theme_fonts": facts["theme_fonts"], "embedded_fonts": facts["embedded"],
        "layouts": [], "slide_list": [],
    }
    for mi, master in enumerate(prs.slide_masters):
        for li, layout in enumerate(master.slide_layouts):
            phs = [f"{p.placeholder_format.idx}:{str(p.placeholder_format.type).split('.')[-1].split(' ')[0]}" for p in layout.placeholders]
            data["layouts"].append({"master": mi, "index": li, "name": layout.name, "placeholders": phs})
    for i, slide in enumerate(prs.slides, start=1):
        if only and i != only:
            continue
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        data["slide_list"].append({"slide": i, "layout": slide.slide_layout.name, "notes": notes[:120],
                                   "shapes": list(_shape_rows(slide.shapes))})
    if as_json:
        print(json.dumps(data, indent=2, ensure_ascii=False))
        return 0
    ratio = sw / sh_
    shape_name = "16:9" if abs(ratio - 16 / 9) < 0.01 else "4:3" if abs(ratio - 4 / 3) < 0.01 else ("A4 portrait" if abs(ratio - 210 / 297) < 0.01 else f"{ratio:.3f}:1")
    print(f"{data['file']} · {data['slides']} slides · {data['size_in'][0]} x {data['size_in'][1]} in ({shape_name}) · {sw} x {sh_} EMU")
    p = data["properties"]
    print(f"properties  title={p['title']!r} author={p['author']!r} last_modified_by={p['last_modified_by']!r}")
    print(f"fonts used  {', '.join(data['fonts_used']) or 'none set on runs'} · theme {', '.join(data['theme_fonts']) or '-'} · embedded {', '.join(data['embedded_fonts']) or 'none'}")
    print("layouts")
    for lay in data["layouts"]:
        print(f"  [{lay['master']}.{lay['index']}] {lay['name']}  placeholders {', '.join(lay['placeholders']) or '-'}")
    for s in data["slide_list"]:
        print(f"\nslide {s['slide']} · layout '{s['layout']}' · notes: {('yes: ' + s['notes'][:60]) if s['notes'] else 'no'}")
        print(f"  {'id':>4}  {'name':<26} {'type':<14} {'x':>6} {'y':>6} {'w':>6} {'h':>6}  text")
        for r in s["shapes"]:
            name = ("  " * r["depth"] + r["name"])[:26]
            nums = " ".join(f"{v:>6.2f}" if v is not None else f"{'-':>6}" for v in (r["x"], r["y"], r["w"], r["h"]))
            text = (r["text"][:70] + "...") if len(r["text"]) > 70 else r["text"]
            print(f"  {r['id']:>4}  {name:<26} {r['type']:<14} {nums}  {text!r}" if text else f"  {r['id']:>4}  {name:<26} {r['type']:<14} {nums}")
    return 0


# ---------------------------------------------------------------- CLI: check (stdlib only)
def _rels(z: zipfile.ZipFile, part: str) -> dict:
    d, f = os.path.split(part)
    rel_path = f"{d}/_rels/{f}.rels"
    if rel_path not in z.namelist():
        return {}
    out = {}
    for rel in ET.fromstring(z.read(rel_path)):
        target = rel.get("Target", "")
        if rel.get("TargetMode") == "External":
            continue
        out[rel.get("Id")] = (rel.get("Type", "").rsplit("/", 1)[-1], os.path.normpath(os.path.join(d, target)).replace("\\", "/"))
    return out


def scan_package(path: str) -> dict:
    """Read the PPTX zip directly: slide order, sizes, extents, fonts, notes. No python-pptx needed."""
    with zipfile.ZipFile(path) as z:
        pres = ET.fromstring(z.read("ppt/presentation.xml"))
        sz = pres.find("p:sldSz", NS)
        slide_w, slide_h = int(sz.get("cx")), int(sz.get("cy"))
        rels = _rels(z, "ppt/presentation.xml")
        order = []
        for sid in pres.findall("p:sldIdLst/p:sldId", NS):
            rid = sid.get(f"{{{NS['r']}}}id")
            if rid in rels:
                order.append(rels[rid][1])
        embedded = sorted({f.get("typeface") for f in pres.findall("p:embeddedFontLst/p:embeddedFont/p:font", NS)})
        theme_fonts = set()
        for name in z.namelist():
            if re.match(r"ppt/theme/theme\d+\.xml$", name):
                t = ET.fromstring(z.read(name))
                for el in t.iter(f"{{{NS['a']}}}latin"):
                    if el.get("typeface"):
                        theme_fonts.add(el.get("typeface"))
        parts = [(f"slide {i}", p) for i, p in enumerate(order, start=1)]
        parts += [(f"layout {os.path.basename(n)}", n) for n in sorted(z.namelist()) if re.match(r"ppt/slideLayouts/slideLayout\d+\.xml$", n)]
        parts += [(f"master {os.path.basename(n)}", n) for n in sorted(z.namelist()) if re.match(r"ppt/slideMasters/slideMaster\d+\.xml$", n)]
        fonts_used, negative, off_slide, empty_ph, notes = set(), [], [], [], {}
        for label, part in parts:
            root = ET.fromstring(z.read(part))
            parent = {c: p for p in root.iter() for c in p}

            def owner(el):
                while el is not None and el.tag not in SHAPE_TAGS:
                    el = parent.get(el)
                if el is None:
                    return "?", "?"
                nv = el.find(".//p:cNvPr", NS)
                return (nv.get("id"), nv.get("name")) if nv is not None else ("?", "?")

            for el in root.iter():
                if el.get("cx") is not None and el.get("cy") is not None:
                    cx_v, cy_v = int(el.get("cx")), int(el.get("cy"))
                    if cx_v < 0 or cy_v < 0:
                        sid, sname = owner(el)
                        negative.append(f"{label}: shape {sid} '{sname}' <{el.tag.split('}')[1]} cx={cx_v} cy={cy_v}>")
                if el.tag == f"{{{NS['a']}}}latin" and el.get("typeface") and not el.get("typeface").startswith("+"):
                    fonts_used.add(el.get("typeface"))
            if not label.startswith("slide"):
                continue
            tree = root.find("p:cSld/p:spTree", NS)
            for sh in list(tree) if tree is not None else []:
                if sh.tag not in SHAPE_TAGS:
                    continue
                nv = sh.find(".//p:cNvPr", NS)
                xfrm = next((x for x in (sh.find("p:spPr/a:xfrm", NS), sh.find("p:grpSpPr/a:xfrm", NS), sh.find("p:xfrm", NS)) if x is not None), None)
                if xfrm is not None and xfrm.find("a:off", NS) is not None and xfrm.find("a:ext", NS) is not None:
                    x, y = int(xfrm.find("a:off", NS).get("x")), int(xfrm.find("a:off", NS).get("y"))
                    w, h = int(xfrm.find("a:ext", NS).get("cx")), int(xfrm.find("a:ext", NS).get("cy"))
                    tol = int(0.02 * EMU_PER_IN)
                    if x < -tol or y < -tol or x + w > slide_w + tol or y + h > slide_h + tol:
                        geo = f"x={x / EMU_PER_IN:.2f} y={y / EMU_PER_IN:.2f} w={w / EMU_PER_IN:.2f} h={h / EMU_PER_IN:.2f} in"
                        off_slide.append((label, nv.get("id"), nv.get("name"), geo))
                if sh.find(".//p:nvPr/p:ph", NS) is not None and sh.tag == f"{{{NS['p']}}}sp":
                    text = "".join(t.text or "" for t in sh.iter(f"{{{NS['a']}}}t")).strip()
                    if not text:
                        empty_ph.append(f"{label}: placeholder {nv.get('id')} '{nv.get('name')}' is empty (edit view shows its prompt)")
            note_text = ""
            for _rid, (kind, target) in _rels(z, part).items():
                if kind == "notesSlide" and target in z.namelist():
                    nroot = ET.fromstring(z.read(target))
                    for sp in nroot.iter(f"{{{NS['p']}}}sp"):
                        ph = sp.find(".//p:nvPr/p:ph", NS)
                        if ph is not None and ph.get("type") == "body":
                            note_text += "".join(t.text or "" for t in sp.iter(f"{{{NS['a']}}}t"))
            notes[label] = note_text.strip()
    return {"slide_w": slide_w, "slide_h": slide_h, "slides": len(order), "negative": negative, "off_slide": off_slide,
            "empty_ph": empty_ph, "notes": notes, "fonts_used": sorted(fonts_used), "theme_fonts": sorted(theme_fonts),
            "embedded": embedded}


def cmd_check(path, max_mb=None, expect_slides=None, require_notes=False, require_embedded=False) -> int:
    try:
        facts = scan_package(path)
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
        print(f"pptx_kit check: FAILED · not a readable .pptx ({exc})", file=sys.stderr)
        return 1
    errors, warns = [], []
    size = os.path.getsize(path)
    errors += [f"NEGATIVE EXTENT  {n}" for n in facts["negative"]]
    grouped: dict = {}
    for label, sid, name, geo in facts["off_slide"]:     # template art repeats on every slide: say it once
        grouped.setdefault((name, geo), []).append((label, sid))
    for (name, geo), where in grouped.items():
        if len(where) == 1:
            warns.append(f"OFF-SLIDE        {where[0][0]}: shape {where[0][1]} '{name}' {geo}")
        else:
            slides = ", ".join(lbl.replace("slide ", "") for lbl, _ in where)
            warns.append(f"OFF-SLIDE        '{name}' {geo} on slides {slides} (the same shape each time: template art bleeding on purpose?)")
    warns += [f"EMPTY PH         {n}" for n in facts["empty_ph"]]
    missing_notes = [k for k, v in facts["notes"].items() if not v]
    if missing_notes:
        (errors if require_notes else warns).append(f"NO NOTES         {', '.join(missing_notes)}")
    if expect_slides is not None and facts["slides"] != expect_slides:
        errors.append(f"SLIDE COUNT      expected {expect_slides}, found {facts['slides']}")
    not_embedded = [f for f in facts["fonts_used"] if f not in facts["embedded"]]
    if require_embedded and (not facts["embedded"] or not_embedded):
        errors.append(f"FONTS            not embedded: {', '.join(not_embedded) or 'none embedded at all'} (re-save from PowerPoint with fonts embedded; see references/powerpoint-com.md)")
    gate = "no size limit set (pass --max-mb)"
    if max_mb is not None:
        if size <= max_mb * 1024 * 1024:
            gate = f"OK <{max_mb:g} MB"
        else:
            gate = f"OVER {max_mb:g} MB"
            errors.append(f"SIZE             {size / 1048576:.1f} MB is over {max_mb:g} MB (JPEG the photos, shrink screenshots to the size they are shown at)")
    w_in, h_in = facts["slide_w"] / EMU_PER_IN, facts["slide_h"] / EMU_PER_IN
    print(f"{os.path.basename(path)} · {facts['slides']} slides · {w_in:.3f} x {h_in:.3f} in · {size / 1048576:.1f} MB · {gate}")
    print(f"  fonts used      {', '.join(facts['fonts_used']) or 'theme fonts only'}")
    print(f"  fonts embedded  {', '.join(facts['embedded']) or 'none'}")
    if not_embedded and not require_embedded:
        print(f"  not embedded    {', '.join(not_embedded)} (fine for system-safe faces; otherwise the viewer substitutes them)")
    for e in errors:
        print(f"  ERROR  {e}")
    for w in warns:
        print(f"  warn   {w}")
    print(f"  result {'FAILED' if errors else 'clean'} · {len(errors)} error(s), {len(warns)} warning(s)")
    return 1 if errors else 0


def main(argv=None) -> int:
    import argparse
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding=None if stream.isatty() else "utf-8", errors="backslashreplace")
        except Exception:
            pass
    ap = argparse.ArgumentParser(prog="pptx_kit.py", description="Inspect a PPTX template; check a finished deck.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    i = sub.add_parser("inspect", help="slide size, layouts, fonts, every shape with id and position")
    i.add_argument("pptx")
    i.add_argument("--slide", type=int, help="only this 1-based slide")
    i.add_argument("--json", action="store_true")
    c = sub.add_parser("check", help="negative extents, off-slide shapes, empty placeholders, notes, fonts, size")
    c.add_argument("pptx")
    c.add_argument("--max-mb", type=float)
    c.add_argument("--expect-slides", type=int)
    c.add_argument("--require-notes", action="store_true")
    c.add_argument("--require-embedded-fonts", action="store_true")
    args = ap.parse_args(argv)
    if not os.path.isfile(args.pptx):
        print(f"pptx_kit: FAILED · not found: {args.pptx}", file=sys.stderr)
        return 1
    if args.cmd == "inspect":
        return cmd_inspect(args.pptx, args.slide, args.json)
    return cmd_check(args.pptx, args.max_mb, args.expect_slides, args.require_notes, args.require_embedded_fonts)


__all__ = ["open_pptx", "emu", "rgb", "shape_by_id", "delete_shape", "send_to_back", "set_notes", "alt_text",
           "clear_text", "replace_text", "text_box", "rect", "line", "card", "chip", "stat", "picture_fit",
           "crop_image", "scan_package", "cmd_inspect", "cmd_check"]

if __name__ == "__main__":
    sys.exit(main())
