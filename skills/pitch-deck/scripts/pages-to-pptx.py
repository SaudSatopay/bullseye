#!/usr/bin/env python3
"""pages-to-pptx.py · build an image-per-slide PPTX from a PDF (route D).

Use it when a platform demands a .pptx but the design must stay identical to the PDF: each slide
is one full-bleed render of a PDF page, so it looks the same on every machine and needs no fonts.
The text is not editable, so every slide gets alt text, and speaker notes when you pass them.

Usage:
    python pages-to-pptx.py <deck.pdf> [--out deck.pptx] [--width 2560] [--format png|jpg]
                            [--quality 90] [--pages 1,3-5] [--notes notes.json] [--alt alt.json]
                            [--title "..."] [--author "..."] [--max-mb 10]

    --width        pixel width of each slide image (default 2560; at least 1920 for projectors)
    --format       png (default, crisp text) or jpg (smaller; use when the PPTX is over a limit)
    --pages        only these 1-based pages, in this order: build an organiser cut from the master PDF
    --notes FILE   JSON array (item 0 = page 1) or object {"1": "...", "2": "..."} keyed by PDF page
    --alt FILE     same shape; default alt text is "Slide N: " + the first words on the page
    --title/--author  document properties (author defaults to empty: no library or machine names)
    --max-mb       exit 1 when the PPTX is larger than this

Needs PyMuPDF and python-pptx:  pip install pymupdf python-pptx
Exit codes: 0 ok · 1 failed or over the limit · 2 bad usage.
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys


def _safe_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            if stream.isatty():
                stream.reconfigure(errors="backslashreplace")
            else:
                stream.reconfigure(encoding="utf-8", errors="backslashreplace")
        except Exception:
            pass


def _need_libs():
    missing = []
    try:
        import fitz  # noqa: F401
    except ImportError:
        missing.append("pymupdf")
    try:
        import pptx  # noqa: F401
    except ImportError:
        missing.append("python-pptx")
    if missing:
        print(f"pages-to-pptx: missing {', '.join(missing)}. Install with:  pip install {' '.join(missing)}", file=sys.stderr)
        sys.exit(1)


def usage(msg: str) -> None:
    print(f"pages-to-pptx: {msg}\nRun with --help for usage.", file=sys.stderr)
    sys.exit(2)


def parse_pages(spec: str, count: int) -> list[int]:
    out: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        try:
            if "-" in part:
                a, b = (int(x) for x in part.split("-", 1))
                out.extend(range(a, b + 1) if a <= b else range(a, b - 1, -1))
            else:
                out.append(int(part))
        except ValueError:
            usage(f"bad page list '{spec}' (use 1,3-5)")
    for p in out:
        if p < 1 or p > count:
            usage(f"page {p} is outside 1-{count}")
    return out


def load_map(path: str | None, what: str) -> dict[int, str]:
    if not path:
        return {}
    try:
        with open(path, encoding="utf-8") as fh:
            data = json.load(fh)
    except Exception as exc:  # noqa: BLE001
        usage(f"cannot read --{what} {path}: {exc}")
    if isinstance(data, list):
        return {i + 1: str(v) for i, v in enumerate(data) if v is not None}
    if isinstance(data, dict):
        out = {}
        for k, v in data.items():
            if not str(k).strip().isdigit():
                usage(f"--{what}: keys must be page numbers, got '{k}'")
            out[int(k)] = str(v)
        return out
    usage(f"--{what} must be a JSON array or an object keyed by page number")
    return {}


def default_alt(page, number: int, total: int) -> str:
    words = re.sub(r"\s+", " ", page.get_text()).strip()
    if not words:
        return f"Slide {number} of {total}"
    if len(words) > 220:
        words = words[:220].rsplit(" ", 1)[0] + " ..."
    return f"Slide {number}: {words}"


def main() -> int:
    _safe_stdio()
    ap = argparse.ArgumentParser(prog="pages-to-pptx.py", description="Image-per-slide PPTX from a PDF, with notes and alt text.")
    ap.add_argument("pdf")
    ap.add_argument("--out", help="output .pptx (default: next to the PDF)")
    ap.add_argument("--width", type=int, default=2560, help="slide image width in pixels (default 2560)")
    ap.add_argument("--format", choices=["png", "jpg"], default="png")
    ap.add_argument("--quality", type=int, default=90, help="JPEG quality for --format jpg (default 90)")
    ap.add_argument("--pages", help="1-based pages to include, in order")
    ap.add_argument("--notes", help="speaker notes JSON")
    ap.add_argument("--alt", help="alt text JSON")
    ap.add_argument("--title")
    ap.add_argument("--author", default="")
    ap.add_argument("--max-mb", type=float)
    args = ap.parse_args()
    if args.width < 800:
        usage("--width must be at least 800 (2560 is the default; 1920 is the floor for projectors)")

    _need_libs()
    import fitz
    from pptx import Presentation
    from pptx.util import Emu

    src = os.path.abspath(args.pdf)
    if not os.path.isfile(src):
        print(f"pages-to-pptx: FAILED · not found: {src}", file=sys.stderr)
        return 1
    out = os.path.abspath(args.out) if args.out else os.path.splitext(src)[0] + ".pptx"
    doc = fitz.open(src)
    if doc.page_count == 0:
        print("pages-to-pptx: FAILED · the PDF has no pages", file=sys.stderr)
        return 1
    pages = parse_pages(args.pages, doc.page_count) if args.pages else list(range(1, doc.page_count + 1))
    notes = load_map(args.notes, "notes")
    alts = load_map(args.alt, "alt")

    first = doc[pages[0] - 1].rect
    prs = Presentation()
    prs.slide_width = Emu(round(first.width * 12700))     # 1 pt = 12,700 EMU
    prs.slide_height = Emu(round(first.height * 12700))
    blank = prs.slide_layouts[6]

    mixed = False
    total = len(pages)
    for i, pno in enumerate(pages, start=1):
        page = doc[pno - 1]
        if abs(page.rect.width / page.rect.height - first.width / first.height) > 0.01:
            mixed = True
        zoom = args.width / page.rect.width
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        data = pix.tobytes("png") if args.format == "png" else pix.tobytes("jpg", jpg_quality=args.quality)
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(io.BytesIO(data), 0, 0, width=prs.slide_width, height=prs.slide_height)
        c_nv_pr = pic._element.nvPicPr.cNvPr
        c_nv_pr.set("name", f"Slide {i} image")
        c_nv_pr.set("descr", alts.get(pno) or default_alt(page, i, total))
        if pno in notes:
            slide.notes_slide.notes_text_frame.text = notes[pno]

    props = prs.core_properties
    props.title = args.title if args.title is not None else ((doc.metadata or {}).get("title") or os.path.splitext(os.path.basename(src))[0])
    props.author = args.author
    props.last_modified_by = args.author
    props.comments = ""
    doc.close()

    os.makedirs(os.path.dirname(out), exist_ok=True)
    try:
        prs.save(out)
    except PermissionError:
        print(f"pages-to-pptx: FAILED · cannot write {out}; is it open in PowerPoint?", file=sys.stderr)
        return 1
    size = os.path.getsize(out)
    noted = sum(1 for p in pages if p in notes)
    gate = "no size limit set (pass --max-mb)"
    status = 0
    if args.max_mb is not None:
        if size <= args.max_mb * 1024 * 1024:
            gate = f"OK <{args.max_mb:g} MB"
        else:
            gate = f"OVER {args.max_mb:g} MB (try --format jpg --quality 85, or --width 1920)"
            status = 1
    print(f"{os.path.basename(out)} · {total} slides · {size / 1048576:.1f} MB · {gate}")
    print(f"  images    {args.width} px wide, {args.format}; slide size {first.width / 72:.3f} x {first.height / 72:.3f} in")
    print(f"  notes     {noted}/{total} slides" + ("" if noted == total else "   <- pass --notes so the presenter has cues (text on these slides is not editable)"))
    print(f"  alt text  {total}/{total} slides" + (" (from --alt)" if alts else " (first words of each page; pass --alt for better descriptions)"))
    if mixed:
        print("  warn      the selected pages do not share one aspect ratio; slides use the first page's size")
    print(f"  wrote     {out}")
    return status


if __name__ == "__main__":
    sys.exit(main())
